"""Excel writer for BOP 40 reports — one sheet per geo, all 5 financial sections."""

from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

from .models import Bop40Report, BopSection, BOP40_GEOS


# ─── Style Constants ───
HEADER_FONT = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="1F4E79")
TOTAL_FONT = Font(name="Calibri", size=11, bold=True)
DATA_FONT = Font(name="Calibri", size=11)
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
SECTION_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
ALERT_FILL = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
THIN_BORDER = Border(
    bottom=Side(style="thin", color="B0B0B0"),
)
CURRENCY_FORMAT = '#,##0.00'


class Bop40ExcelWriter:
    """Generates a multi-sheet Excel workbook for BOP 40 reports."""

    def generate(self, reports: dict[str, Bop40Report], output_path: str) -> str:
        wb = Workbook()
        wb.remove(wb.active)  # remove default sheet

        # Summary sheet
        self._write_summary_sheet(wb, reports)

        # One sheet per geo
        for geo in BOP40_GEOS:
            if geo in reports:
                self._write_geo_sheet(wb, reports[geo])

        wb.save(output_path)
        return output_path

    def _write_summary_sheet(self, wb: Workbook, reports: dict[str, Bop40Report]):
        ws = wb.create_sheet("Summary")
        ws.column_dimensions["A"].width = 25
        for col in range(2, 8):
            ws.column_dimensions[get_column_letter(col)].width = 18

        # Header row
        row = 1
        ws.cell(row=row, column=1, value="BOP 40 Report — Cross-Geo Summary").font = HEADER_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        for col in range(1, 8):
            ws.cell(row=row, column=col).fill = HEADER_FILL

        row = 3
        headers = ["Metric", "US (USD)", "EU (EUR)", "JP (JPY)", "IN (INR)", "BR (BRL)"]
        for i, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=i, value=h)
            cell.font = SECTION_FONT
            cell.fill = SECTION_FILL

        metrics = [
            ("Total Assets", "total_assets"),
            ("Total Liabilities", "total_liabilities"),
            ("Total Equity", "total_equity"),
            ("Net Income", "net_income"),
            ("BS Check (A-L-E)", "balance_sheet_check"),
        ]

        for metric_name, attr in metrics:
            row += 1
            ws.cell(row=row, column=1, value=metric_name).font = DATA_FONT
            for i, geo in enumerate(BOP40_GEOS, 2):
                if geo in reports:
                    val = float(getattr(reports[geo], attr))
                    cell = ws.cell(row=row, column=i, value=val)
                    cell.number_format = CURRENCY_FORMAT
                    cell.font = DATA_FONT
                    if attr == "balance_sheet_check" and abs(val) > 0.01:
                        cell.fill = ALERT_FILL

    def _write_geo_sheet(self, wb: Workbook, report: Bop40Report):
        ws = wb.create_sheet(f"{report.geo_code} ({report.currency})")
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 35
        ws.column_dimensions["C"].width = 18

        row = 1
        title = f"BOP 40 — {report.geo_code} | {report.period} | {report.currency}"
        ws.cell(row=row, column=1, value=title).font = HEADER_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        for col in range(1, 4):
            ws.cell(row=row, column=col).fill = HEADER_FILL

        # Write each section
        row = 3
        for section in [report.assets, report.liabilities, report.equity,
                        report.revenue, report.expenses]:
            row = self._write_section(ws, section, row)
            row += 1  # blank row between sections

        # Balance sheet check
        row += 1
        ws.cell(row=row, column=1, value="").font = TOTAL_FONT
        ws.cell(row=row, column=2, value="Balance Sheet Check (A - L - E)").font = TOTAL_FONT
        check_cell = ws.cell(row=row, column=3, value=float(report.balance_sheet_check))
        check_cell.number_format = CURRENCY_FORMAT
        check_cell.font = TOTAL_FONT
        if abs(float(report.balance_sheet_check)) > 0.01:
            check_cell.fill = ALERT_FILL
        else:
            check_cell.fill = TOTAL_FILL

    def _write_section(self, ws, section: BopSection, start_row: int) -> int:
        row = start_row

        # Section header
        ws.cell(row=row, column=1, value="").fill = SECTION_FILL
        ws.cell(row=row, column=2, value=section.section_name).font = SECTION_FONT
        ws.cell(row=row, column=2).fill = SECTION_FILL
        ws.cell(row=row, column=3, value="Amount").font = SECTION_FONT
        ws.cell(row=row, column=3).fill = SECTION_FILL
        row += 1

        for item in section.line_items:
            indent = "  " * item.rollup_level
            ws.cell(row=row, column=1, value=item.account_id).font = DATA_FONT
            name_cell = ws.cell(row=row, column=2, value=f"{indent}{item.account_name}")
            amt_cell = ws.cell(row=row, column=3, value=float(item.amount))
            amt_cell.number_format = CURRENCY_FORMAT

            if item.is_rollup_total:
                name_cell.font = TOTAL_FONT
                amt_cell.font = TOTAL_FONT
                amt_cell.fill = TOTAL_FILL
                for col in range(1, 4):
                    ws.cell(row=row, column=col).border = Border(
                        top=Side(style="thin"), bottom=Side(style="double"))
            else:
                name_cell.font = DATA_FONT
                amt_cell.font = DATA_FONT
                amt_cell.border = THIN_BORDER

            row += 1

        return row
